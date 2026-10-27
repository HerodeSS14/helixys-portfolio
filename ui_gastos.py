import os
from datetime import datetime
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit, QPushButton,
    QHeaderView, QTreeWidget, QTreeWidgetItem, QTabWidget, QMessageBox, 
    QDoubleSpinBox, QDateEdit, QDialog, QFrame, QFileDialog, QComboBox, QListWidget, QListWidgetItem
)
from PySide6.QtCore import Qt, QDate, QSignalBlocker
from PySide6.QtGui import QColor, QBrush, QKeyEvent

from database import SessionLocal
from models import Gasto
from estilos import tema

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# =============================================================================
# WIDGET PERSONALIZADO PARA ENTRADA ESTILO ATM (DE DERECHA A IZQUIERDA)
# =============================================================================
class SpinBoxMontoATM(QDoubleSpinBox):
    """
    QDoubleSpinBox personalizado que permite la entrada de números
    desplazándose de derecha a izquierda (estilo punto de venta / centavos).
    """
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setRange(0, 100000000)
        self.setDecimals(2)
        self.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        self.setGroupSeparatorShown(True)
        self._buffer_digitos = ""

    def focusInEvent(self, event):
        super().focusInEvent(event)
        self.selectAll()
        if self.value() > 0:
            self._buffer_digitos = f"{int(round(self.value() * 100))}"
        else:
            self._buffer_digitos = ""

    def focusOutEvent(self, event):
        super().focusOutEvent(event)
        self._buffer_digitos = ""

    def keyPressEvent(self, event: QKeyEvent):
        key = event.key()
        
        if Qt.Key_0 <= key <= Qt.Key_9:
            digit = chr(key)
            if self._buffer_digitos == "" and digit == "0":
                return
            
            if len(self._buffer_digitos) < 10:
                self._buffer_digitos += digit
                val = int(self._buffer_digitos) / 100.0
                self.setValue(val)
            return

        elif key == Qt.Key_Backspace:
            if len(self._buffer_digitos) > 0:
                self._buffer_digitos = self._buffer_digitos[:-1]
                val = int(self._buffer_digitos) / 100.0 if self._buffer_digitos else 0.0
                self.setValue(val)
            else:
                self.setValue(0.0)
            return

        super().keyPressEvent(event)


# =============================================================================
# DIÁLOGO PARA ADMINISTRAR EL CATÁLOGO DE GASTOS
# =============================================================================
class DialogoAdministrarCatalogo(QDialog):
    """ Ventana Modal para Crear o Eliminar ítems del Catálogo de Gastos """
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("⚙️ Administrar Catálogo de Gastos")
        self.setMinimumSize(420, 480)
        self.init_ui()
        self.cargar_catalogo()

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(12)

        lbl_titulo = QLabel("<b>Catálogo de Gastos / Servicios</b>")
        lbl_titulo.setStyleSheet(tema.ESTILO_LABEL_TITULO)
        layout.addWidget(lbl_titulo)

        lbl_sub = QLabel("Agregue nuevos tipos de gastos o elimine los que ya no utilice:")
        lbl_sub.setStyleSheet("font-size: 12px; color: #666;")
        layout.addWidget(lbl_sub)

        # Entrada para nuevo elemento
        h_add = QHBoxLayout()
        self.in_nuevo_gasto = QLineEdit()
        self.in_nuevo_gasto.setPlaceholderText("Ej: Condominio, Patente, Seguro...")
        self.in_nuevo_gasto.setMinimumHeight(38)
        self.in_nuevo_gasto.setStyleSheet(tema.ESTILO_ENTRY_MODERNO)

        self.btn_crear = QPushButton("➕ Crear")
        self.btn_crear.setMinimumHeight(38)
        self.btn_crear.setStyleSheet(tema.ESTILO_BTN_PRIMARIO)
        self.btn_crear.setCursor(Qt.PointingHandCursor)
        self.btn_crear.clicked.connect(self.crear_gasto)

        h_add.addWidget(self.in_nuevo_gasto)
        h_add.addWidget(self.btn_crear)
        layout.addLayout(h_add)

        # Lista de gastos existentes
        self.lista_gastos = QListWidget()
        self.lista_gastos.setStyleSheet(f"""
            QListWidget {{
                background-color: {tema.BG_INPUT};
                border: 1px solid {tema.BORDE_SUAVE};
                border-radius: 6px;
                padding: 5px;
                font-size: 13px;
            }}
            QListWidget::item {{
                padding: 8px;
                border-bottom: 1px solid #E2E8F0;
            }}
            QListWidget::item:selected {{
                background-color: #3498DB;
                color: white;
            }}
        """)
        layout.addWidget(self.lista_gastos)

        # Botones inferiores
        h_bottom = QHBoxLayout()
        self.btn_eliminar_cat = QPushButton("🗑️ Eliminar Seleccionado")
        self.btn_eliminar_cat.setMinimumHeight(38)
        self.btn_eliminar_cat.setStyleSheet(tema.ESTILO_BTN_PELIGRO)
        self.btn_eliminar_cat.setCursor(Qt.PointingHandCursor)
        self.btn_eliminar_cat.clicked.connect(self.eliminar_gasto_catalogo)

        self.btn_cerrar = QPushButton("Cerrar")
        self.btn_cerrar.setMinimumHeight(38)
        self.btn_cerrar.setStyleSheet(tema.ESTILO_BTN_SECUNDARIO)
        self.btn_cerrar.setCursor(Qt.PointingHandCursor)
        self.btn_cerrar.clicked.connect(self.accept)

        h_bottom.addWidget(self.btn_eliminar_cat)
        h_bottom.addWidget(self.btn_cerrar)
        layout.addLayout(h_bottom)

    def cargar_catalogo(self):
        self.lista_gastos.clear()
        db = SessionLocal()
        try:
            # Consulta únicamente lo registrado en la BD sin listas duras por defecto
            registrados = db.query(Gasto.concepto).distinct().all()
            gastos_registrados = [r[0].strip() for r in registrados if r[0] and r[0].strip()]
        finally:
            db.close()

        for g in sorted(gastos_registrados):
            self.lista_gastos.addItem(g)

    def crear_gasto(self):
        nombre = self.in_nuevo_gasto.text().strip()
        if not nombre:
            QMessageBox.warning(self, "Aviso", "Escriba el nombre del nuevo gasto.", QMessageBox.Ok)
            return

        items = [self.lista_gastos.item(i).text().lower() for i in range(self.lista_gastos.count())]
        if nombre.lower() in items:
            QMessageBox.warning(self, "Aviso", "Este gasto ya existe en el catálogo.", QMessageBox.Ok)
            return

        db = SessionLocal()
        try:
            nuevo = Gasto(concepto=nombre, monto=0.0, moneda="USD", fecha=datetime.now())
            db.add(nuevo)
            db.commit()
            self.in_nuevo_gasto.clear()
            self.cargar_catalogo()
            QMessageBox.information(self, "Éxito", f"Gasto '{nombre}' agregado al catálogo.", QMessageBox.Ok)
        except Exception as e:
            db.rollback()
            QMessageBox.critical(self, "Error", f"No se pudo crear: {e}", QMessageBox.Ok)
        finally:
            db.close()

    def eliminar_gasto_catalogo(self):
        item_sel = self.lista_gastos.currentItem()
        if not item_sel:
            QMessageBox.warning(self, "Aviso", "Seleccione un gasto de la lista para eliminar.", QMessageBox.Ok)
            return

        concepto = item_sel.text()
        box = QMessageBox(
            QMessageBox.Question,
            "Confirmar Eliminación",
            f"¿Desea eliminar '{concepto}' y todos sus registros asociados?",
            QMessageBox.Yes | QMessageBox.No,
            self
        )
        if box.exec() == QMessageBox.Yes:
            db = SessionLocal()
            try:
                db.query(Gasto).filter(Gasto.concepto == concepto).delete()
                db.commit()
                self.cargar_catalogo()
                QMessageBox.information(self, "Eliminado", f"Gasto '{concepto}' eliminado correctamente.", QMessageBox.Ok)
            except Exception as e:
                db.rollback()
                QMessageBox.critical(self, "Error", f"Error al eliminar: {e}", QMessageBox.Ok)
            finally:
                db.close()


# =============================================================================
# PANTALLA PRINCIPAL DE GASTOS
# =============================================================================
class PantallaGastos(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.init_ui()
        tema.tema_cambiado.connect(self.actualizar_diseno_dinamico)
        self.actualizar_diseno_dinamico()
        
        self.tabs.currentChanged.connect(self.al_cambiar_pestana)

        self.cargar_gastos_mes_actual()
        self.cargar_historial_gastos()

    def init_ui(self):
        layout_principal = QVBoxLayout(self)
        layout_principal.setContentsMargins(15, 15, 15, 15)
        layout_principal.setSpacing(15)

        self.tabs = QTabWidget()

        # ------------------- TAB 1: GASTOS DEL MES ACTUAL -------------------
        self.tab_mes_actual = QWidget()
        lay_tab_mes = QVBoxLayout(self.tab_mes_actual)
        lay_tab_mes.setContentsMargins(10, 10, 10, 10)
        lay_tab_mes.setSpacing(10)

        # Encabezado dinámico
        self.lbl_titulo_mes = QLabel("💸 Gastos del Mes en Curso")
        self.lbl_titulo_mes.setStyleSheet(tema.ESTILO_LABEL_TITULO)
        lay_tab_mes.addWidget(self.lbl_titulo_mes)

        lbl_leyenda = QLabel("💡 <b>Edición Directa ATM:</b> Digite los números en la celda de monto de derecha a izquierda (ej: digite '1000' para 10,00).")
        lbl_leyenda.setStyleSheet("font-size: 12px; color: #444;")
        lay_tab_mes.addWidget(lbl_leyenda)

        self.arbol_mes_actual = QTreeWidget()
        self.arbol_mes_actual.setHeaderLabels([
            "Gasto / Servicio", "Moneda", "Monto Gastado", "Fecha del Gasto"
        ])
        lay_tab_mes.addWidget(self.arbol_mes_actual)

        # Tarjeta de Totales
        self.card_totales = QFrame()
        self.card_totales.setFrameShape(QFrame.StyledPanel)
        lay_totales = QHBoxLayout(self.card_totales)
        lay_totales.setContentsMargins(15, 10, 15, 10)

        self.lbl_total_usd = QLabel("<b>TOTAL USD:</b> $ 0.00")
        self.lbl_total_usd.setStyleSheet("font-size: 14px; font-weight: bold; color: #E74C3C;")
        
        self.lbl_total_ves = QLabel("<b>TOTAL VES:</b> Bs. 0.00")
        self.lbl_total_ves.setStyleSheet("font-size: 14px; font-weight: bold; color: #27AE60;")

        lay_totales.addWidget(self.lbl_total_usd)
        lay_totales.addSpacing(30)
        lay_totales.addWidget(self.lbl_total_ves)
        lay_totales.addStretch()

        lay_tab_mes.addWidget(self.card_totales)

        # Botón para Administrar Catálogo
        h_acciones = QHBoxLayout()
        self.btn_admin_cat = QPushButton("⚙️ Administrar Catálogo")
        self.btn_admin_cat.setMinimumHeight(40)
        self.btn_admin_cat.setCursor(Qt.PointingHandCursor)
        self.btn_admin_cat.clicked.connect(self.abrir_administrar_catalogo)
        h_acciones.addWidget(self.btn_admin_cat)
        h_acciones.addStretch()

        lay_tab_mes.addLayout(h_acciones)
        self.tabs.addTab(self.tab_mes_actual, "📝 Gastos del Mes")

        # ------------------- TAB 2: HISTORIAL DE GASTOS -------------------
        self.tab_historial = QWidget()
        lay_tab_hist = QVBoxLayout(self.tab_historial)
        lay_tab_hist.setContentsMargins(15, 15, 15, 15)
        lay_tab_hist.setSpacing(15)

        h_top_hist = QHBoxLayout()
        self.lbl_hist_titulo = QLabel("📊 Historial Anual y Mensual de Gastos")
        self.lbl_hist_titulo.setStyleSheet(tema.ESTILO_LABEL_TITULO)
        h_top_hist.addWidget(self.lbl_hist_titulo)

        self.btn_exportar_excel = QPushButton("📥 EXPORTAR MES A EXCEL")
        self.btn_exportar_excel.setMinimumHeight(40)
        self.btn_exportar_excel.setCursor(Qt.PointingHandCursor)
        self.btn_exportar_excel.clicked.connect(self.exportar_mes_excel)
        h_top_hist.addWidget(self.btn_exportar_excel)

        lay_tab_hist.addLayout(h_top_hist)

        self.arbol_historial = QTreeWidget()
        self.arbol_historial.setHeaderLabels([
            "Período / Gasto", "Fecha", "Moneda", "Monto Total"
        ])
        lay_tab_hist.addWidget(self.arbol_historial)

        self.tabs.addTab(self.tab_historial, "📊 Historial por Mes / Año")
        layout_principal.addWidget(self.tabs)

    def al_cambiar_pestana(self, index):
        if index == 1:
            self.cargar_historial_gastos()

    def actualizar_diseno_dinamico(self):
        color_bg_input = tema.BG_INPUT
        self.setStyleSheet(tema.ESTILO_VENTANA)

        style_tree = f"""
            QTreeWidget {{
                background-color: {color_bg_input};
                color: #000000;
                font-size: 13px;
                border: 1px solid {tema.BORDE_SUAVE};
                border-radius: 6px;
            }}
            QHeaderView::section {{
                background-color: #93F6E3;
                color: #1C1C1C;
                font-size: 13px;
                font-weight: bold;
                padding: 10px 8px;
                border-right: 1px solid #CBD5E1;
                border-bottom: 2px solid #CBD5E1;
            }}
            QTreeWidget::item {{
                min-height: 44px;
                padding: 4px 8px;
                border-bottom: 1px solid #CBD5E1;
                color: #000000;
            }}
        """
        for tree in [self.arbol_mes_actual, self.arbol_historial]:
            tree.setStyleSheet(style_tree)

        header_hist = self.arbol_historial.header()
        header_hist.setSectionResizeMode(0, QHeaderView.Stretch)
        header_hist.setSectionResizeMode(1, QHeaderView.ResizeToContents)
        header_hist.setSectionResizeMode(2, QHeaderView.ResizeToContents)
        header_hist.setSectionResizeMode(3, QHeaderView.ResizeToContents)

        self.btn_admin_cat.setStyleSheet(tema.ESTILO_BTN_PRIMARIO)
        self.btn_exportar_excel.setStyleSheet("QPushButton { background-color: #27ae60; color: white; font-weight: bold; border-radius: 6px; padding: 8px 16px; border: none; } QPushButton:hover { background-color: #219150; }")

    # =========================================================================
    # CARGA DE GASTOS DEL MES ACTUAL CON EDICIÓN DIRECTA EN CELDAS
    # =========================================================================
    def cargar_gastos_mes_actual(self):
        self.arbol_mes_actual.clear()
        db = SessionLocal()
        try:
            ahora = datetime.now()
            nombres_meses = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
            self.lbl_titulo_mes.setText(f"💸 Gastos de {nombres_meses[ahora.month - 1]} {ahora.year}")

            # 1. Obtener catálogo únicamente guardado en BD
            registrados = db.query(Gasto.concepto).distinct().all()
            gastos_base = sorted(list(set([r[0].strip() for r in registrados if r[0] and r[0].strip()])))

            # 2. Obtener gastos guardados del mes en curso
            todos_gastos = db.query(Gasto).all()
            gastos_mes = {
                g.concepto.strip(): g for g in todos_gastos 
                if g.fecha and g.fecha.year == ahora.year and g.fecha.month == ahora.month
            }

            for concepto in gastos_base:
                item = QTreeWidgetItem(self.arbol_mes_actual, [concepto, "", "", ""])
                gasto_existente = gastos_mes.get(concepto)

                g_id = gasto_existente.id if gasto_existente else None
                monto_val = gasto_existente.monto if gasto_existente else 0.0
                moneda_val = gasto_existente.moneda if gasto_existente else "USD"
                
                if gasto_existente and gasto_existente.fecha:
                    fecha_val = QDate(gasto_existente.fecha.year, gasto_existente.fecha.month, gasto_existente.fecha.day)
                else:
                    fecha_val = QDate.currentDate()

                item.setData(0, Qt.UserRole, {"concepto": concepto, "id": g_id})

                # --- CONTROL DE MONEDA EN CELDA ---
                cmb_moneda = QComboBox()
                cmb_moneda.addItems(["USD ($)", "VES (Bs.)"])
                cmb_moneda.setMinimumHeight(32)
                cmb_moneda.setStyleSheet(tema.ESTILO_ENTRY_MODERNO)
                if moneda_val == "VES":
                    cmb_moneda.setCurrentIndex(1)
                else:
                    cmb_moneda.setCurrentIndex(0)

                # --- CONTROL DE MONTO EN CELDA ESTILO ATM ---
                spn_monto = SpinBoxMontoATM()
                spn_monto.setValue(monto_val)
                spn_monto.setMinimumHeight(32)
                spn_monto.setPrefix("Bs. " if moneda_val == "VES" else "$ ")
                spn_monto.setStyleSheet(tema.ESTILO_ENTRY_MODERNO)

                # --- CONTROL DE FECHA EN CELDA ---
                in_fecha = QDateEdit()
                in_fecha.setCalendarPopup(True)
                in_fecha.setDate(fecha_val)
                in_fecha.setMinimumHeight(32)
                in_fecha.setStyleSheet(tema.ESTILO_ENTRY_MODERNO)

                spn_monto.editingFinished.connect(
                    lambda item=item, cmb=cmb_moneda, spn=spn_monto, fch=in_fecha: self._guardar_cambio_directo(item, cmb, spn, fch)
                )

                cmb_moneda.currentIndexChanged.connect(
                    lambda idx, item=item, cmb=cmb_moneda, spn=spn_monto, fch=in_fecha: self._guardar_cambio_directo(item, cmb, spn, fch)
                )
                
                in_fecha.dateChanged.connect(
                    lambda qdt, item=item, cmb=cmb_moneda, spn=spn_monto, fch=in_fecha: self._guardar_cambio_directo(item, cmb, spn, fch)
                )

                self.arbol_mes_actual.setItemWidget(item, 1, cmb_moneda)
                self.arbol_mes_actual.setItemWidget(item, 2, spn_monto)
                self.arbol_mes_actual.setItemWidget(item, 3, in_fecha)

            self.recalcular_totales_mes()
            self.arbol_mes_actual.resizeColumnToContents(0)
            self.arbol_mes_actual.setColumnWidth(1, 130)
            self.arbol_mes_actual.setColumnWidth(2, 180)
            self.arbol_mes_actual.setColumnWidth(3, 160)

        finally:
            db.close()

    # =========================================================================
    # GUARDADO AUTOMÁTICO AL EDITAR EN LA TABLA
    # =========================================================================
    def _guardar_cambio_directo(self, item, cmb_moneda, spn_monto, in_fecha):
        data = item.data(0, Qt.UserRole)
        if not data: return

        concepto = data["concepto"]
        g_id = data.get("id")
        moneda = "VES" if cmb_moneda.currentIndex() == 1 else "USD"
        spn_monto.setPrefix("Bs. " if moneda == "VES" else "$ ")
        monto = spn_monto.value()

        qdt = in_fecha.date()
        fecha_dt = datetime(qdt.year(), qdt.month(), qdt.day())

        db = SessionLocal()
        try:
            gasto_mes = None
            if g_id:
                gasto_mes = db.query(Gasto).filter(Gasto.id == g_id).first()

            if not gasto_mes:
                gasto_mes = db.query(Gasto).filter(
                    Gasto.concepto == concepto,
                    Gasto.fecha == fecha_dt
                ).first()

            if not gasto_mes:
                if monto > 0:
                    gasto_mes = Gasto(concepto=concepto, categoria="General", fecha=fecha_dt, moneda=moneda, monto=monto)
                    db.add(gasto_mes)
            else:
                gasto_mes.moneda = moneda
                gasto_mes.monto = monto
                gasto_mes.fecha = fecha_dt

            db.commit()
            if gasto_mes:
                data["id"] = gasto_mes.id
                item.setData(0, Qt.UserRole, data)

            self.recalcular_totales_mes()
        except Exception as e:
            db.rollback()
            print(f"Error al guardar cambio directo de gasto: {e}")
        finally:
            db.close()

    def recalcular_totales_mes(self):
        total_usd = 0.0
        total_ves = 0.0

        root = self.arbol_mes_actual.invisibleRootItem()
        for i in range(root.childCount()):
            item = root.child(i)
            cmb = self.arbol_mes_actual.itemWidget(item, 1)
            spn = self.arbol_mes_actual.itemWidget(item, 2)
            if cmb and spn:
                monto = spn.value()
                if cmb.currentIndex() == 1:
                    total_ves += monto
                else:
                    total_usd += monto

        self.lbl_total_usd.setText(f"<b>TOTAL GASTADO (USD):</b> $ {total_usd:,.2f}")
        self.lbl_total_ves.setText(f"<b>TOTAL GASTADO (VES):</b> Bs. {total_ves:,.2f}")

    # =========================================================================
    # CARGA DEL HISTORIAL JERÁRQUICO
    # =========================================================================
    def cargar_historial_gastos(self):
        self.arbol_historial.clear()
        db = SessionLocal()
        try:
            gastos = db.query(Gasto).order_by(Gasto.fecha.desc()).all()

            estructura = {}
            nombres_meses = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

            for g in gastos:
                if not g.fecha or g.monto <= 0: continue
                ano = g.fecha.year
                mes_num = g.fecha.month

                if ano not in estructura: estructura[ano] = {}
                if mes_num not in estructura[ano]: estructura[ano][mes_num] = []

                estructura[ano][mes_num].append(g)

            for ano in sorted(estructura.keys(), reverse=True):
                item_ano = QTreeWidgetItem(self.arbol_historial, [f"📅 AÑO {ano}"])
                font_ano = item_ano.font(0)
                font_ano.setBold(True); font_ano.setPointSize(12)
                item_ano.setFont(0, font_ano)
                item_ano.setForeground(0, QBrush(QColor("#3498db")))

                total_ano_usd = 0.0
                total_ano_ves = 0.0

                for mes_num in sorted(estructura[ano].keys(), reverse=True):
                    lista_m = estructura[ano][mes_num]
                    nombre_mes = nombres_meses[mes_num - 1]

                    total_mes_usd = sum(g.monto for g in lista_m if getattr(g, 'moneda', 'USD') == 'USD')
                    total_mes_ves = sum(g.monto for g in lista_m if getattr(g, 'moneda', 'USD') == 'VES')

                    total_ano_usd += total_mes_usd
                    total_ano_ves += total_mes_ves

                    tot_mes_str = f"TOTAL: $ {total_mes_usd:,.2f}"
                    if total_mes_ves > 0:
                        tot_mes_str += f" | Bs. {total_mes_ves:,.2f}"

                    item_mes = QTreeWidgetItem(item_ano, [
                        f"🗓️ {nombre_mes} {ano} ({len(lista_m)} gastos)", "", "", tot_mes_str
                    ])
                    item_mes.setData(0, Qt.UserRole, {"tipo": "MES", "ano": ano, "mes": mes_num, "nombre_mes": nombre_mes})
                    font_mes = item_mes.font(0)
                    font_mes.setBold(True)
                    item_mes.setFont(0, font_mes)
                    item_mes.setForeground(0, QBrush(QColor("#2ecc71")))

                    for g in lista_m:
                        f_str = g.fecha.strftime("%d/%m/%Y")
                        prefix = "Bs. " if g.moneda == "VES" else "$ "
                        QTreeWidgetItem(item_mes, [
                            g.concepto, f_str, g.moneda, f"{prefix}{g.monto:,.2f}"
                        ])

                tot_ano_str = f"TOTAL AÑO: $ {total_ano_usd:,.2f}"
                if total_ano_ves > 0:
                    tot_ano_str += f" | Bs. {total_ano_ves:,.2f}"
                item_ano.setText(3, tot_ano_str)
                item_ano.setExpanded(True)

            header_hist = self.arbol_historial.header()
            header_hist.setSectionResizeMode(0, QHeaderView.Stretch)
            header_hist.setSectionResizeMode(1, QHeaderView.ResizeToContents)
            header_hist.setSectionResizeMode(2, QHeaderView.ResizeToContents)
            header_hist.setSectionResizeMode(3, QHeaderView.ResizeToContents)

        finally:
            db.close()

    # =========================================================================
    # ACCIONES
    # =========================================================================
    def abrir_administrar_catalogo(self):
        dlg = DialogoAdministrarCatalogo(self)
        if dlg.exec() == QDialog.Accepted:
            self.cargar_gastos_mes_actual()
            self.cargar_historial_gastos()

    def exportar_mes_excel(self):
        if not HAS_OPENPYXL:
            QMessageBox.critical(self, "Librería Faltante", "La librería 'openpyxl' no está instalada.\nEjecute: pip install openpyxl", QMessageBox.Ok)
            return

        item_sel = self.arbol_historial.currentItem()
        if not item_sel:
            QMessageBox.warning(self, "Aviso", "Seleccione un MES en el árbol del historial para exportar.", QMessageBox.Ok)
            return

        data_role = item_sel.data(0, Qt.UserRole)
        if not data_role or data_role.get("tipo") != "MES":
            QMessageBox.warning(self, "Aviso", "Haga clic sobre el encabezado del mes que desea exportar.", QMessageBox.Ok)
            return

        ano = data_role["ano"]
        mes_num = data_role["mes"]
        nombre_mes = data_role["nombre_mes"]

        file_path, _ = QFileDialog.getSaveFileName(self, f"Exportar Gastos {nombre_mes} {ano}", f"Gastos_{nombre_mes}_{ano}.xlsx", "Archivos Excel (*.xlsx)")
        if not file_path: return

        db = SessionLocal()
        try:
            gastos = db.query(Gasto).all()
            filtrados = [g for g in gastos if g.fecha and g.fecha.year == ano and g.fecha.month == mes_num and g.monto > 0]

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = f"Gastos {nombre_mes} {ano}"

            font_titulo = Font(name="Arial", size=13, bold=True, color="FFFFFF")
            fill_titulo = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
            
            font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            fill_header = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
            
            font_total = Font(name="Arial", size=11, bold=True, color="15803D")
            border_thin = Border(left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'), top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))

            ws.merge_cells("A1:D1")
            cell_t = ws.cell(row=1, column=1, value=f"REPORTE DE GASTOS - {nombre_mes.upper()} {ano}")
            cell_t.font = font_titulo; cell_t.fill = fill_titulo
            cell_t.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 35

            headers = ["Fecha", "Gasto / Servicio", "Moneda", "Monto Total"]
            ws.append([])
            ws.append(headers)
            ws.row_dimensions[3].height = 25

            for col_idx in range(1, 5):
                cell = ws.cell(row=3, column=col_idx)
                cell.font = font_header; cell.fill = fill_header
                cell.alignment = Alignment(horizontal="center", vertical="center")

            total_usd = 0.0
            total_ves = 0.0

            for idx, g in enumerate(filtrados, start=4):
                f_str = g.fecha.strftime("%d/%m/%Y")
                moneda_str = getattr(g, 'moneda', 'USD')

                if moneda_str == "VES": total_ves += g.monto
                else: total_usd += g.monto

                ws.append([f_str, g.concepto, moneda_str, g.monto])
                cell_monto = ws.cell(row=idx, column=4)
                cell_monto.number_format = '"$"#,##0.00' if moneda_str == "USD" else '"Bs."#,##0.00'

                for c_i in range(1, 5):
                    ws.cell(row=idx, column=c_i).border = border_thin

            row_tot = len(filtrados) + 5
            ws.cell(row=row_tot, column=3, value="TOTAL GASTADO ($):").font = font_total
            ws.cell(row=row_tot, column=3).alignment = Alignment(horizontal="right")
            c_usd = ws.cell(row=row_tot, column=4, value=total_usd)
            c_usd.font = font_total
            c_usd.number_format = '"$"#,##0.00'

            if total_ves > 0:
                row_tot_ves = row_tot + 1
                ws.cell(row=row_tot_ves, column=3, value="TOTAL GASTADO (Bs.):").font = font_total
                ws.cell(row=row_tot_ves, column=3).alignment = Alignment(horizontal="right")
                c_ves = ws.cell(row=row_tot_ves, column=4, value=total_ves)
                c_ves.font = font_total
                c_ves.number_format = '"Bs."#,##0.00'

            for col in ws.columns:
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = 22

            wb.save(file_path)
            QMessageBox.information(self, "Exportación Exitosa", f"Archivo Excel generado con éxito en:\n{file_path}", QMessageBox.Ok)

        except Exception as e:
            QMessageBox.critical(self, "Error al Exportar", f"No se pudo generar el archivo Excel: {e}", QMessageBox.Ok)
        finally:
            db.close()