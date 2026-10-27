import os
from datetime import datetime

# --- LIBRERÍAS REPORTLAB (PDF) ---
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import inch, mm
from reportlab.lib.utils import ImageReader, simpleSplit
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Image, Table, TableStyle, KeepTogether, PageBreak

# --- LIBRERÍAS PYTHON-DOCX (WORD EDITABLE) ---
try:
    import docx
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    DOCX_DISPONIBLE = True
except ImportError:
    DOCX_DISPONIBLE = False


class MotorImpresion:
    CARPETA_RECIPES = "recipes_generados"

    @staticmethod
    def _obtener_ruta_unica(ruta_base, nombre_base, extension="pdf"):
        """Verifica si el archivo existe. Si existe, agrega _1, _2, etc."""
        ext = extension.lstrip('.')
        ruta_completa = os.path.join(ruta_base, f"{nombre_base}.{ext}")
        contador = 1
        while os.path.exists(ruta_completa):
            ruta_completa = os.path.join(ruta_base, f"{nombre_base}_{contador}.{ext}")
            contador += 1
        return ruta_completa

    @staticmethod
    def _asegurar_carpeta():
        if not os.path.exists(MotorImpresion.CARPETA_RECIPES):
            os.makedirs(MotorImpresion.CARPETA_RECIPES)

    @staticmethod
    def _ajustar_extension_salida(ruta_salida, formato="pdf"):
        """Asegura que la ruta de salida coincida con el formato solicitado (pdf o docx)."""
        base, _ = os.path.splitext(ruta_salida)
        ext = ".docx" if formato.lower() == "docx" else ".pdf"
        return f"{base}{ext}"

    # =========================================================================
    # --- MÉTODOS AUXILIARES REPORTLAB (PDF) ---
    # =========================================================================

    @staticmethod
    def _dibujar_cabecera(c, doctor, paciente, titulo_doc):
        """Dibuja exclusivamente el membrete superior oficial calculando las coordenadas para A4"""
        ancho_hoja, alto_hoja = A4  # 210mm x 297mm
        
        es_fem = hasattr(doctor, 'genero_doctor') and doctor.genero_doctor and doctor.genero_doctor.strip().lower() in ["femenino", "mujer"]
        prefijo = "Dra." if es_fem else "Dr."
        
        # 1. Logo
        if doctor.logo_ruta and os.path.exists(doctor.logo_ruta):
            logo = ImageReader(doctor.logo_ruta)
            c.drawImage(logo, 0.5*inch, alto_hoja - 1.4*inch, width=1.2*inch, height=1.1*inch, mask='auto')

        # 2. Datos del Doctor
        c.setFont("Helvetica-Bold", 14)
        c.drawRightString(ancho_hoja - 0.5*inch, alto_hoja - 0.6*inch, doctor.nombre_clinica.upper())
        c.setFont("Helvetica-Bold", 12)
        c.drawRightString(ancho_hoja - 0.5*inch, alto_hoja - 0.8*inch, f"{prefijo} {doctor.nombre_doctor}")
        c.setFont("Helvetica", 10)
        c.drawRightString(ancho_hoja - 0.5*inch, alto_hoja - 0.95*inch, doctor.especialidades)
        c.drawRightString(ancho_hoja - 0.5*inch, alto_hoja - 1.1*inch, f"C.I: {doctor.cedula_doctor} | MSDS: {doctor.matricula_msds}")
        
        c.setLineWidth(1.5)
        c.line(0.5*inch, alto_hoja - 1.5*inch, ancho_hoja - 0.5*inch, alto_hoja - 1.5*inch)

        # 3. Ficha del Paciente
        c.setFont("Helvetica-Bold", 11)
        c.drawString(0.5*inch, alto_hoja - 1.8*inch, f"PACIENTE: {paciente.nombre_completo}")
        c.drawRightString(ancho_hoja - 0.5*inch, alto_hoja - 1.8*inch, f"FECHA: {datetime.now().strftime('%d/%m/%Y')}")
        
        c.setFont("Helvetica", 10)
        tiene_seguro = hasattr(paciente, 'seguro') and paciente.seguro and paciente.seguro.strip().lower() not in ["", "no", "ninguno", "n/a", "no aplica"]
        
        if tiene_seguro:
            c.drawString(0.5*inch, alto_hoja - 2.0*inch, f"CÉDULA: {paciente.cedula}   |   EDAD: {paciente.edad_actual} años   |   SEGURO: {paciente.seguro.upper()}")
        else:
            c.drawString(0.5*inch, alto_hoja - 2.0*inch, f"CÉDULA: {paciente.cedula}   |   EDAD: {paciente.edad_actual} años")
        
        c.setLineWidth(1)
        c.line(0.5*inch, alto_hoja - 2.1*inch, ancho_hoja - 0.5*inch, alto_hoja - 2.1*inch)
        
        c.setFont("Helvetica-Bold", 14)
        c.drawCentredString(ancho_hoja/2, alto_hoja - 2.4*inch, titulo_doc.upper())

    @staticmethod
    def _dibujar_pie_estatico_decorativo(c, doctor):
        """Dibuja únicamente la información de contacto en la parte inferior de la página"""
        ancho_hoja, alto_hoja = A4
        c.setFont("Helvetica-Oblique", 8)
        c.setStrokeColorRGB(0.7, 0.7, 0.7)
        c.setLineWidth(0.5)
        c.line(0.5*inch, 0.6*inch, ancho_hoja - 0.5*inch, 0.6*inch)
        c.drawCentredString(ancho_hoja/2, 0.45*inch, f"Dirección: {doctor.direccion_clinica}")
        c.drawCentredString(ancho_hoja/2, 0.3*inch, f"Rif: {doctor.rif_clinica} | {doctor.correo_doctor}")

    @staticmethod
    def _obtener_bloque_firma_fluyente(doctor):
        """Construye una estructura de tabla con la firma para ser inyectada directamente en el flujo dinámico"""
        es_fem = hasattr(doctor, 'genero_doctor') and doctor.genero_doctor and doctor.genero_doctor.strip().lower() in ["femenino", "mujer"]
        prefijo = "Dra." if es_fem else "Dr."
        
        style_lbl = ParagraphStyle('FirmaMedLbl', fontName='Helvetica', fontSize=10, leading=12, alignment=1)
        celda_elementos = []
        
        if doctor.firma_ruta and os.path.exists(doctor.firma_ruta):
            img_firma = Image(doctor.firma_ruta, width=2.2*inch, height=0.9*inch)
            img_firma.hAlign = 'CENTER'
            celda_elementos.append(img_firma)
        else:
            celda_elementos.append(Spacer(1, 45))
            
        celda_elementos.append(Paragraph("_______________________________________", style_lbl))
        celda_elementos.append(Paragraph(f"<b>{prefijo} {doctor.nombre_doctor}</b>", style_lbl))
        celda_elementos.append(Paragraph(f"Médico Especialista | MSDS: {doctor.matricula_msds}", style_lbl))
        
        t_firma = Table([[celda_elementos]], colWidths=[4.0*inch])
        t_firma.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE')
        ]))
        t_firma.hAlign = 'CENTER'
        return t_firma

    # =========================================================================
    # --- MÉTODOS AUXILIARES DOCX (WORD EDITABLE) ---
    # =========================================================================

    @staticmethod
    def _crear_documento_docx():
        """Inicializa un documento Word A4 con márgenes apropiados."""
        if not DOCX_DISPONIBLE:
            raise ImportError("La librería 'python-docx' no está instalada. Ejecuta 'pip install python-docx'.")
        
        doc = Document()
        for section in doc.sections:
            section.page_width = Inches(8.27)
            section.page_height = Inches(11.69)
            section.top_margin = Inches(0.6)
            section.bottom_margin = Inches(0.6)
            section.left_margin = Inches(0.6)
            section.right_margin = Inches(0.6)
        return doc

    @staticmethod
    def _agregar_cabecera_docx(doc, doctor, paciente, titulo_doc):
        """Genera el membrete médico equivalente en Word."""
        es_fem = hasattr(doctor, 'genero_doctor') and doctor.genero_doctor and doctor.genero_doctor.strip().lower() in ["femenino", "mujer"]
        prefijo = "Dra." if es_fem else "Dr."

        tbl_top = doc.add_table(rows=1, cols=2)
        tbl_top.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_logo, cell_info = tbl_top.rows[0].cells
        cell_logo.width = Inches(2.0)
        cell_info.width = Inches(5.0)

        if doctor.logo_ruta and os.path.exists(doctor.logo_ruta):
            p_logo = cell_logo.paragraphs[0]
            p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_logo.add_run().add_picture(doctor.logo_ruta, width=Inches(1.2))

        p_info = cell_info.paragraphs[0]
        p_info.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        run_clinica = p_info.add_run(f"{doctor.nombre_clinica.upper()}\n")
        run_clinica.bold = True
        run_clinica.font.size = Pt(13)
        run_clinica.font.color.rgb = RGBColor(41, 128, 185)

        run_doc = p_info.add_run(f"{prefijo} {doctor.nombre_doctor}\n")
        run_doc.bold = True
        run_doc.font.size = Pt(11)

        run_esp = p_info.add_run(f"{doctor.especialidades}\n")
        run_esp.font.size = Pt(9.5)

        run_ids = p_info.add_run(f"C.I: {doctor.cedula_doctor} | MSDS: {doctor.matricula_msds}")
        run_ids.font.size = Pt(9)

        p_div = doc.add_paragraph()
        p_div.paragraph_format.space_before = Pt(4)
        p_div.paragraph_format.space_after = Pt(6)
        run_line = p_div.add_run("_________________________________________________________________________________")
        run_line.font.color.rgb = RGBColor(180, 180, 180)

        p_pac = doc.add_paragraph()
        p_pac.paragraph_format.space_before = Pt(2)
        p_pac.paragraph_format.space_after = Pt(2)

        nombre_p = paciente.get("nombre_completo", "") if isinstance(paciente, dict) else getattr(paciente, "nombre_completo", "")
        cedula_p = paciente.get("cedula", "") if isinstance(paciente, dict) else getattr(paciente, "cedula", "")
        edad_p = paciente.get("edad_actual", "--") if isinstance(paciente, dict) else getattr(paciente, "edad_actual", "--")

        run_p_lbl = p_pac.add_run("PACIENTE: ")
        run_p_lbl.bold = True
        run_p_val = p_pac.add_run(f"{nombre_p.upper()}\t\t")
        run_p_val.bold = True

        run_f_lbl = p_pac.add_run("FECHA: ")
        run_f_lbl.bold = True
        p_pac.add_run(f"{datetime.now().strftime('%d/%m/%Y')}")

        p_det = doc.add_paragraph()
        p_det.paragraph_format.space_after = Pt(6)
        
        tiene_seguro = hasattr(paciente, 'seguro') and paciente.seguro and paciente.seguro.strip().lower() not in ["", "no", "ninguno", "n/a", "no aplica"]
        txt_seguro = f"   |   SEGURO: {paciente.seguro.upper()}" if tiene_seguro else ""
        p_det.add_run(f"CÉDULA: {cedula_p}   |   EDAD: {edad_p} años{txt_seguro}")

        p_tit = doc.add_paragraph()
        p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_tit.paragraph_format.space_before = Pt(8)
        p_tit.paragraph_format.space_after = Pt(12)
        run_tit = p_tit.add_run(titulo_doc.upper())
        run_tit.bold = True
        run_tit.font.size = Pt(13)
        run_tit.font.color.rgb = RGBColor(30, 41, 59)

    @staticmethod
    def _agregar_firma_docx(doc, doctor):
        """Agrega la estructura de firma en Word."""
        es_fem = hasattr(doctor, 'genero_doctor') and doctor.genero_doctor and doctor.genero_doctor.strip().lower() in ["femenino", "mujer"]
        prefijo = "Dra." if es_fem else "Dr."

        p_firma = doc.add_paragraph()
        p_firma.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_firma.paragraph_format.space_before = Pt(24)
        p_firma.paragraph_format.space_after = Pt(0)

        if doctor.firma_ruta and os.path.exists(doctor.firma_ruta):
            p_firma.add_run().add_picture(doctor.firma_ruta, height=Inches(0.8))
            p_firma.add_run("\n")

        run_linea = p_firma.add_run("_______________________________________\n")
        run_linea.bold = True

        run_doc = p_firma.add_run(f"{prefijo} {doctor.nombre_doctor}\n")
        run_doc.bold = True
        run_doc.font.size = Pt(10.5)

        run_sub = p_firma.add_run(f"Médico Especialista | MSDS: {doctor.matricula_msds}")
        run_sub.font.size = Pt(9)

    @staticmethod
    def _agregar_pie_pagina_docx(doc, doctor):
        """Agrega el pie de página en Word."""
        section = doc.sections[0]
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_ft = p_ft.add_run(f"Dirección: {doctor.direccion_clinica}\nRif: {doctor.rif_clinica} | {doctor.correo_doctor}")
        run_ft.font.size = Pt(8)
        run_ft.font.italic = True
        run_ft.font.color.rgb = RGBColor(120, 120, 120)

    # =========================================================================
    # --- MÉTODOS PÚBLICOS PRINCIPALES CON SOPORTE PDF / DOCX ---
    # =========================================================================

    @staticmethod
    def generar_pdf_recipe(doctor, paciente, lista_medicamentos, formato="pdf"):
        """Genera récipe e indicaciones. Soporta formato="pdf" o formato="docx"."""
        if formato.lower() == "docx":
            ruta_base = os.path.abspath("recipes_generados")
            os.makedirs(ruta_base, exist_ok=True)
            nombre_base = f"Recipe_{paciente.cedula}_{datetime.now().strftime('%Y-%m-%d')}"
            ruta_docx = MotorImpresion._obtener_ruta_unica(ruta_base, nombre_base, extension="docx")

            doc = MotorImpresion._crear_documento_docx()

            # Página 1: Récipe
            MotorImpresion._agregar_cabecera_docx(doc, doctor, paciente, "RÉCIPE MÉDICO")
            p_rp = doc.add_paragraph()
            p_rp.paragraph_format.space_before = Pt(8)
            p_rp.paragraph_format.space_after = Pt(8)
            run_rp = p_rp.add_run("Rp:")
            run_rp.bold = True
            run_rp.font.size = Pt(14)

            for med, _ in lista_medicamentos:
                p_m = doc.add_paragraph()
                p_m.paragraph_format.space_after = Pt(4)
                run_m = p_m.add_run(f"• {med.strip().title()}")
                run_m.bold = True
                run_m.font.size = Pt(11)

            MotorImpresion._agregar_firma_docx(doc, doctor)

            # Salto de página para Indicaciones
            doc.add_page_break()

            # Página 2: Indicaciones
            MotorImpresion._agregar_cabecera_docx(doc, doctor, paciente, "INDICACIONES MÉDICAS")
            for med, ind in lista_medicamentos:
                p_med = doc.add_paragraph()
                p_med.paragraph_format.space_before = Pt(6)
                p_med.paragraph_format.space_after = Pt(2)
                run_m = p_med.add_run(f"Medicamento: {med.strip().title()}")
                run_m.bold = True
                run_m.font.size = Pt(11)

                p_ind = doc.add_paragraph()
                p_ind.paragraph_format.left_indent = Inches(0.3)
                p_ind.paragraph_format.space_after = Pt(10)
                p_ind.add_run(f"Indicaciones: {ind.strip()}")

            MotorImpresion._agregar_firma_docx(doc, doctor)
            MotorImpresion._agregar_pie_pagina_docx(doc, doctor)

            doc.save(ruta_docx)
            return os.path.abspath(ruta_docx)

        # --- GENERACIÓN PDF ORIGINAL ---
        ruta_base = os.path.abspath("recipes_generados")
        os.makedirs(ruta_base, exist_ok=True)
        
        nombre_base = f"Recipe_{paciente.cedula}_{datetime.now().strftime('%Y-%m-%d')}"
        ruta_completa = MotorImpresion._obtener_ruta_unica(ruta_base, nombre_base, extension="pdf")
        
        doc = SimpleDocTemplate(
            ruta_completa, 
            pagesize=A4, 
            leftMargin=0.5*inch, 
            rightMargin=0.5*inch,
            topMargin=2.5*inch,    
            bottomMargin=0.9*inch  
        )
        
        styles = getSampleStyleSheet()
        estilo_rp = ParagraphStyle('RpTxt', fontName='Helvetica-Bold', fontSize=16, leading=18, spaceBefore=10, spaceAfter=10)
        estilo_med = ParagraphStyle('MedTxt', fontName='Helvetica-Bold', fontSize=11, leading=14, spaceBefore=6)
        estilo_ind = ParagraphStyle('IndTxt', fontName='Helvetica', fontSize=10, leading=14, leftIndent=15, spaceAfter=12, alignment=4)
        
        elementos = []
        elementos.append(Paragraph("Rp:", estilo_rp))
        for med, _ in lista_medicamentos:
            med_formateado = med.strip().title()
            elementos.append(Paragraph(f"• {med_formateado}", estilo_med))
            
        elementos.append(Spacer(1, 20))
        elementos.append(MotorImpresion._obtener_bloque_firma_fluyente(doctor))
        elementos.append(PageBreak())
        
        for med, ind in lista_medicamentos:
            bloque_med = []
            med_formateado = med.strip().title()
            bloque_med.append(Paragraph(f"Medicamento: {med_formateado}", estilo_med))
            bloque_med.append(Paragraph(f"Indicaciones: {ind.strip()}", estilo_ind))
            elementos.append(KeepTogether(bloque_med))
            
        elementos.append(Spacer(1, 20))
        elementos.append(KeepTogether([MotorImpresion._obtener_bloque_firma_fluyente(doctor)]))

        def callback_paginas(canvas_obj, doc_obj):
            titulo = "RÉCIPE MÉDICO" if doc_obj.page == 1 else "INDICACIONES MÉDICAS"
            MotorImpresion._dibujar_cabecera(canvas_obj, doctor, paciente, titulo)
            MotorImpresion._dibujar_pie_estatico_decorativo(canvas_obj, doctor)

        try:
            doc.build(elementos, onFirstPage=callback_paginas, onLaterPages=callback_paginas)
            return os.path.abspath(ruta_completa)
        except Exception as e:
            raise Exception(f"Error generando el Récipe Integrado A4: {str(e)}")

    @staticmethod
    def generar_pdf_informe(doctor, paciente, texto_informe, ruta_salida, formato="pdf"):
        """Genera un Informe Médico. Soporta formato="pdf" o formato="docx"."""
        if formato.lower() == "docx":
            ruta_docx = MotorImpresion._ajustar_extension_salida(ruta_salida, "docx")
            doc = MotorImpresion._crear_documento_docx()
            MotorImpresion._agregar_cabecera_docx(doc, doctor, paciente, "INFORME MÉDICO")

            for parrafo in texto_informe.split('\n'):
                if parrafo.strip():
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.line_spacing = 1.15
                    run = p.add_run(parrafo.strip())
                    run.font.size = Pt(11)

            MotorImpresion._agregar_firma_docx(doc, doctor)
            MotorImpresion._agregar_pie_pagina_docx(doc, doctor)
            doc.save(ruta_docx)
            return os.path.abspath(ruta_docx)

        # --- GENERACIÓN PDF ORIGINAL ---
        ruta_salida_pdf = MotorImpresion._ajustar_extension_salida(ruta_salida, "pdf")
        doc = SimpleDocTemplate(ruta_salida_pdf, pagesize=A4, rightMargin=50, leftMargin=50, topMargin=2.5*inch, bottomMargin=2.0*inch)
        styles = getSampleStyleSheet()
        elementos = []
        estilo_cuerpo = ParagraphStyle('CuerpoInforme', parent=styles['Normal'], fontSize=11, leading=15, alignment=4)
        
        elementos.append(Spacer(1, 20))
        texto_formateado = texto_informe.replace('\n', '<br/>')
        elementos.append(Paragraph(texto_formateado, estilo_cuerpo))
        
        elementos.append(Spacer(1, 30))
        elementos.append(KeepTogether([MotorImpresion._obtener_bloque_firma_fluyente(doctor)]))

        def construir_pagina(canvas, doc_obj):
            MotorImpresion._dibujar_cabecera(canvas, doctor, paciente, "INFORME MÉDICO")
            MotorImpresion._dibujar_pie_estatico_decorativo(canvas, doctor)
            canvas.setFont("Helvetica", 8)

        try:
            doc.build(elementos, onFirstPage=construir_pagina, onLaterPages=construir_pagina)
            return os.path.abspath(ruta_salida_pdf)
        except Exception as e:
            raise Exception(f"Error al generar el PDF del informe: {str(e)}")

    @staticmethod
    def generar_pdf_historia(doctor, paciente, consulta, ruta_salida, formato="pdf"):
        """Genera la Historia Clínica. Soporta formato="pdf" o formato="docx"."""
        def es_valido(valor):
            return valor is not None and str(valor).strip().lower() not in ["no aplica", "none", "", "false", "0", "0.0"]

        if formato.lower() == "docx":
            ruta_docx = MotorImpresion._ajustar_extension_salida(ruta_salida, "docx")
            doc = MotorImpresion._crear_documento_docx()
            titulo_hist = f"HISTORIA CLÍNICA - {consulta.fecha_consulta.strftime('%d/%m/%Y')}"
            MotorImpresion._agregar_cabecera_docx(doc, doctor, paciente, titulo_hist)

            def add_sec(tit):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                r = p.add_run(tit)
                r.bold = True
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(41, 128, 185)

            def add_txt(lbl, val):
                if es_valido(val):
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(3)
                    r_lbl = p.add_run(f"{lbl}: ")
                    r_lbl.bold = True
                    p.add_run(str(val))

            add_sec("I. MOTIVO DE CONSULTA")
            add_txt("Motivo de Consulta", getattr(consulta, 'motivo_consulta', None))
            add_txt("Enfermedad Actual", consulta.enfermedad_actual)
            add_txt("Evolución", consulta.nota_evolutiva)

            add_sec("II. ANTECEDENTES Y HÁBITOS")
            ant = paciente.antecedentes
            if ant:
                add_txt("Alergias", ant.alergias)
                add_txt("Personales / Patológicos", ant.personales)
                add_txt("Familiares", ant.familiares)
                add_txt("Operaciones Realizadas", ant.operaciones_realizadas)
                
                habitos = []
                for label, val in [("Fuma", ant.habito_fuma), ("Bebe", ant.habito_bebe), ("Otros", ant.otros_habitos)]:
                    if es_valido(val): habitos.append(f"{label}: {val}")
                if habitos:
                    add_txt("Hábitos", " | ".join(habitos))
                add_txt("Examen Funcional", ant.examen_funcional)

            add_sec("III. SIGNOS VITALES Y BIOMETRÍA")
            p_sv = doc.add_paragraph()
            sv_arr = []
            if es_valido(consulta.peso): sv_arr.append(f"Peso: {consulta.peso} kg")
            if es_valido(consulta.altura): sv_arr.append(f"Altura: {consulta.altura} m")
            if es_valido(consulta.imc): sv_arr.append(f"IMC: {consulta.imc:.2f}")
            if es_valido(consulta.tension_arterial): sv_arr.append(f"Tensión: {consulta.tension_arterial}")
            if es_valido(consulta.pulso): sv_arr.append(f"Pulso: {consulta.pulso} ppm")
            p_sv.add_run("  |  ".join(sv_arr) if sv_arr else "Sin registros.")

            add_sec("IV. EXAMEN FÍSICO")
            campos_gen = [("Piel", consulta.piel), ("Ojos", consulta.ojos), ("Oídos", consulta.oidos),
                          ("Orofaringe", consulta.orofaringe), ("Tiroides", consulta.tiroides),
                          ("Axilas", consulta.axilas), ("Supraclaviculares", consulta.supraclaviculares),
                          ("Abdomen", consulta.abdomen), ("Lumbar", consulta.lumbar),
                          ("Neurológico", getattr(consulta, 'neurologico', None)),
                          ("Otros Hallazgos", getattr(consulta, 'otros_examen_fisico', None))]
            for lbl, val in campos_gen:
                add_txt(lbl, val)

            add_sec("V. DIAGNÓSTICO Y PLAN DE ESTUDIO")
            add_txt("Impresión Diagnóstica", getattr(consulta, 'impresion_diagnostica', None))
            add_txt("Diagnóstico", getattr(consulta, 'diagnostico', None))
            add_txt("Estudios / Plan Indicado", getattr(consulta, 'estudios_indicados', None))

            MotorImpresion._agregar_firma_docx(doc, doctor)
            MotorImpresion._agregar_pie_pagina_docx(doc, doctor)
            doc.save(ruta_docx)
            return os.path.abspath(ruta_docx)

        # --- GENERACIÓN PDF ORIGINAL ---
        ruta_salida_pdf = MotorImpresion._ajustar_extension_salida(ruta_salida, "pdf")
        ancho_hoja, alto_hoja = A4
        doc = SimpleDocTemplate(ruta_salida_pdf, pagesize=A4, topMargin=2.4*inch, bottomMargin=1*inch, leftMargin=0.5*inch, rightMargin=0.5*inch)
        styles = getSampleStyleSheet()
        estilo_seccion = ParagraphStyle('Seccion', parent=styles['Heading3'], color=colors.HexColor("#2980b9"), spaceBefore=12, spaceAfter=6)
        estilo_texto = ParagraphStyle('Texto', parent=styles['Normal'], fontSize=10, leading=14)
        
        elementos = []

        elementos.append(Paragraph("<b>I. MOTIVO DE CONSULTA</b>", estilo_seccion))
        if es_valido(getattr(consulta, 'motivo_consulta', None)):
            elementos.append(Paragraph(f"<b>Motivo de Consulta:</b> {consulta.motivo_consulta}", estilo_texto))
        elementos.append(Paragraph(f"<b>Enfermedad Actual:</b> {consulta.enfermedad_actual}", estilo_texto))
        if es_valido(consulta.nota_evolutiva):
            elementos.append(Paragraph(f"<b>Evolución:</b> {consulta.nota_evolutiva}", estilo_texto))

        elementos.append(Paragraph("<b>II. ANTECEDENTES Y HÁBITOS</b>", estilo_seccion))
        ant = paciente.antecedentes
        if ant:
            if es_valido(ant.alergias): elementos.append(Paragraph(f"<b>Alergias:</b> {ant.alergias}", estilo_texto))
            if es_valido(ant.personales): elementos.append(Paragraph(f"<b>Personales / Patológicos:</b> {ant.personales}", estilo_texto))
            if es_valido(ant.familiares): elementos.append(Paragraph(f"<b>Familiares:</b> {ant.familiares}", estilo_texto))
            if es_valido(ant.operaciones_realizadas): elementos.append(Paragraph(f"<b>Operaciones Realizadas:</b> {ant.operaciones_realizadas}", estilo_texto))
            
            habitos = []
            for label, val in [("Fuma", ant.habito_fuma), ("Bebe", ant.habito_bebe), ("Otros", ant.otros_habitos)]:
                if es_valido(val): habitos.append(f"<b>{label}:</b> {val}")
            if habitos:
                elementos.append(Paragraph(f"<b>Hábitos:</b> " + " | ".join(habitos), estilo_texto))
                
            if es_valido(ant.examen_funcional):
                elementos.append(Paragraph(f"<b>Examen Funcional:</b> {ant.examen_funcional}", estilo_texto))
        else:
            elementos.append(Paragraph("Sin antecedentes registrados.", estilo_texto))

        elementos.append(Paragraph("<b>III. SIGNOS VITALES Y BIOMETRÍA</b>", estilo_seccion))
        sv_data = []
        if es_valido(consulta.peso) or es_valido(consulta.altura):
            sv_data.append(["Peso:", f"{consulta.peso} kg" if es_valido(consulta.peso) else "-",
                            "Altura:", f"{consulta.altura} m" if es_valido(consulta.altura) else "-",
                            "IMC:", f"{consulta.imc:.2f}" if es_valido(consulta.imc) else "-"])
            
        sv_data.append(["Tensión:", str(consulta.tension_arterial) if consulta.tension_arterial else "-",
                        "Pulso:", f"{consulta.pulso} ppm" if es_valido(consulta.pulso) else "-", "", ""])
        
        ancho_util_pulgadas = (ancho_hoja / inch) - 1.0
        t_sv = Table(sv_data, colWidths=[ancho_util_pulgadas/6 * inch]*6)
        t_sv.setStyle(TableStyle([
            ('GRID', (0,0), (-1,-1), 0.5, colors.lightgrey),
            ('BACKGROUND', (0,0), (0,-1), colors.HexColor("#f9f9f9")),
            ('BACKGROUND', (2,0), (2,-1), colors.HexColor("#f9f9f9")),
            ('BACKGROUND', (4,0), (4,-1), colors.HexColor("#f9f9f9")),
            ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
            ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
            ('FONTNAME', (2,0), (2,-1), 'Helvetica-Bold'),
            ('FONTNAME', (4,0), (4,-1), 'Helvetica-Bold'),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ]))
        elementos.append(t_sv)

        elementos.append(Paragraph("<b>IV. EXAMEN FÍSICO </b>", estilo_seccion))
        ef_gen_data = []
        campos_gen = [("Piel", consulta.piel), ("Ojos", consulta.ojos), ("Oídos", consulta.oidos),
                      ("Orofaringe", consulta.orofaringe), ("Tiroides", consulta.tiroides),
                      ("Axilas", consulta.axilas), ("Supraclaviculares", consulta.supraclaviculares),
                      ("Abdomen", consulta.abdomen), ("Lumbar", consulta.lumbar),
                      ("Neurológico", getattr(consulta, 'neurologico', None)),
                      ("Otros Hallazgos", getattr(consulta, 'otros_examen_fisico', None))]
        
        for label, valor in campos_gen:
            if es_valido(valor):
                ef_gen_data.append([Paragraph(f"<b>{label}:</b>", estilo_texto), Paragraph(str(valor), estilo_texto)])
        
        if ef_gen_data:
            t_ef = Table(ef_gen_data, colWidths=[1.5*inch, (ancho_util_pulgadas-1.5)*inch])
            t_ef.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP'), ('BOTTOMPADDING', (0,0), (-1,-1), 4)]))
            elementos.append(t_ef)
        else:
            elementos.append(Paragraph("Sin hallazgos patológicos registrados.", estilo_texto))

        elementos.append(Paragraph("<b>V. EXAMEN ESPECIALIZADO</b>", estilo_seccion))
        esp_data = []
        if paciente.genero == "Femenino":
            campos_esp = [("Mamas", consulta.mamas), ("Vulva", consulta.vulva), ("Vagina", consulta.vagina),
                          ("Cérvix", consulta.cervix), ("Colposcopia", consulta.colposcopia), 
                          ("Tacto Recto-Vaginal", consulta.tacto_recto_vaginal)]
        else:
            campos_esp = [("Pene/Testículos", consulta.pene_testiculos), ("Próstata", consulta.prostata)]
            
        for label, valor in campos_esp:
            if es_valido(valor):
                esp_data.append([Paragraph(f"<b>{label}:</b>", estilo_texto), Paragraph(str(valor), estilo_texto)])
                
        if esp_data:
            t_esp = Table(esp_data, colWidths=[1.5*inch, (ancho_util_pulgadas-1.5)*inch])
            t_esp.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP'), ('BOTTOMPADDING', (0,0), (-1,-1), 4)]))
            elementos.append(t_esp)
        else:
            elementos.append(Paragraph("No evaluado / Sin hallazgos.", estilo_texto))

        es_mujer = paciente.genero and paciente.genero.strip().lower() in ["femenino", "mujer"]
        if es_mujer:
            from database import SessionLocal
            from models import GinecoObstetrico
            db = SessionLocal()
            try:
                gin = db.query(GinecoObstetrico).filter(GinecoObstetrico.paciente_id == paciente.id).order_by(GinecoObstetrico.fecha_registro.desc()).first()
                if gin:
                    elementos.append(Paragraph("<b>VI. ANTECEDENTES GINECO-OBSTÉTRICOS</b>", estilo_seccion))
                    incont = "Sí" if gin.incontinencia else "No"
                    campos_gineco = [
                        ("Menarquia", gin.menarquia), ("Gestas", gin.gestas), ("Abortos", gin.abortos),
                        ("Emb. Ectópico", gin.embarazo_ectopico), ("Partos", gin.partos), ("Cesáreas", gin.cesarea), 
                        ("Menopausia", gin.menopausia), ("Sangramiento", gin.sangramiento), ("Incontinencia", incont),
                        ("Flujo", gin.flujo), ("DIU", gin.diu), ("ACO", gin.aco)
                    ]
                    
                    gin_arr = []
                    for label, val in campos_gineco:
                        if label in ["Gestas", "Abortos", "Emb. Ectópico", "Partos", "Cesáreas", "Incontinencia"]:
                            if val is not None:
                                gin_arr.append(f"<b>{label}:</b> {val}")
                        elif es_valido(val): 
                            gin_arr.append(f"<b>{label}:</b> {val}")
                    
                    if gin_arr:
                        elementos.append(Paragraph(" | ".join(gin_arr), estilo_texto))
            finally:
                db.close()

        elementos.append(Paragraph("<b>VII. DIAGNÓSTICO Y PLAN DE ESTUDIO</b>", estilo_seccion))
        if es_valido(getattr(consulta, 'impresion_diagnostica', None)):
            elementos.append(Paragraph(f"<b>Impresión Diagnóstica:</b> {consulta.impresion_diagnostica}", estilo_texto))
        if es_valido(getattr(consulta, 'diagnostico', None)):
            elementos.append(Paragraph(f"<b>Diagnóstico:</b> {consulta.diagnostico}", estilo_texto))
        if es_valido(getattr(consulta, 'estudios_indicados', None)):
            elementos.append(Paragraph(f"<b>Estudios / Plan Indicado:</b> {consulta.estudios_indicados}", estilo_texto))

        elementos.append(Spacer(1, 20))
        elementos.append(KeepTogether([MotorImpresion._obtener_bloque_firma_fluyente(doctor)]))

        def footer_historia(canvas, doc_obj):
            MotorImpresion._dibujar_cabecera(canvas, doctor, paciente, f"HISTORIA CLÍNICA - {consulta.fecha_consulta.strftime('%d/%m/%Y')}")
            MotorImpresion._dibujar_pie_estatico_decorativo(canvas, doctor)
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor(colors.grey)

        doc.build(elementos, onFirstPage=footer_historia, onLaterPages=footer_historia)
        return ruta_salida_pdf

    @staticmethod
    def generar_pdf_consentimiento(doctor, paciente, consentimiento, plantilla_base, ruta_salida, formato="pdf"):
        """Genera Consentimiento Informado. Soporta formato="pdf" o formato="docx"."""
        es_fem = hasattr(doctor, 'genero_doctor') and doctor.genero_doctor and doctor.genero_doctor.strip().lower() in ["femenino", "mujer"]
        prefijo = "Dra." if es_fem else "Dr."

        if formato.lower() == "docx":
            ruta_docx = MotorImpresion._ajustar_extension_salida(ruta_salida, "docx")
            doc = MotorImpresion._crear_documento_docx()
            MotorImpresion._agregar_cabecera_docx(doc, doctor, paciente, "CONSENTIMIENTO INFORMADO PARA CIRUGÍA")

            texto = plantilla_base
            texto = texto.replace("[PACIENTE]", f"{paciente.nombre_completo.upper()}")
            texto = texto.replace("[CEDULA_PACIENTE]", f"{paciente.cedula}")
            texto = texto.replace("[DOCTOR]", f"{prefijo} {doctor.nombre_doctor.upper()}")
            texto = texto.replace("[CEDULA_DOCTOR]", f"{doctor.cedula_doctor}")
            texto = texto.replace("[ESPECIALIDAD]", f"{doctor.especialidades}")
            texto = texto.replace("[MSDS]", f"{doctor.matricula_msds}")
            texto = texto.replace("[INTERVENCION]", f"{consentimiento.intervencion.upper()}")
            texto = texto.replace("[DIAGNOSTICO]", f"{consentimiento.diagnostico.upper()}")
            texto = texto.replace("[RIESGOS]", f"{consentimiento.riesgos_especificos.upper()}")
            
            txt_testigo = f"\nActuando en representación de mi familiar {consentimiento.familiar_nombre} (C.I. {consentimiento.familiar_cedula})." if consentimiento.familiar_nombre else ""
            texto = texto.replace("[BLOQUE_TESTIGO]", txt_testigo)

            for p_txt in texto.split('\n'):
                if p_txt.strip():
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.line_spacing = 1.15
                    run = p.add_run(p_txt.strip())
                    run.font.size = Pt(10)

            p_f = doc.add_paragraph()
            p_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_f.paragraph_format.space_before = Pt(20)
            p_f.add_run("_____________________________________\nFIRMA DEL PACIENTE\n" + f"{paciente.nombre_completo}\nC.I: {paciente.cedula}")

            if consentimiento.familiar_nombre:
                p_f.add_run("\n\n_____________________________________\nFIRMA DEL FAMILIAR / TESTIGO\n" + f"{consentimiento.familiar_nombre}\nC.I: {consentimiento.familiar_cedula}")

            MotorImpresion._agregar_firma_docx(doc, doctor)
            MotorImpresion._agregar_pie_pagina_docx(doc, doctor)
            doc.save(ruta_docx)
            return os.path.abspath(ruta_docx)

        # --- GENERACIÓN PDF ORIGINAL ---
        ruta_salida_pdf = MotorImpresion._ajustar_extension_salida(ruta_salida, "pdf")
        ancho_hoja, alto_hoja = A4
        doc = SimpleDocTemplate(ruta_salida_pdf, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=2.4*inch, bottomMargin=1.4*inch)
        styles = getSampleStyleSheet()
        estilo_texto = ParagraphStyle('TextoLegal', parent=styles['Normal'], fontSize=10, leading=14, alignment=4, spaceAfter=8)
        
        elementos = []
        texto = plantilla_base
        texto = texto.replace("[PACIENTE]", f"<b>{paciente.nombre_completo.upper()}</b>")
        texto = texto.replace("[CEDULA_PACIENTE]", f"<b>{paciente.cedula}</b>")
        texto = texto.replace("[DOCTOR]", f"<b>{prefijo} {doctor.nombre_doctor.upper()}</b>")
        texto = texto.replace("[CEDULA_DOCTOR]", f"<b>{doctor.cedula_doctor}</b>")
        texto = texto.replace("[ESPECIALIDAD]", f"<b>{doctor.especialidades}</b>")
        texto = texto.replace("[MSDS]", f"<b>{doctor.matricula_msds}</b>")
        texto = texto.replace("[INTERVENCION]", f"<b>{consentimiento.intervencion.upper()}</b>")
        texto = texto.replace("[DIAGNOSTICO]", f"<b>{consentimiento.diagnostico.upper()}</b>")
        texto = texto.replace("[RIESGOS]", f"<b>{consentimiento.riesgos_especificos.upper()}</b>")
        
        if consentimiento.familiar_nombre:
            txt_testigo = f"<br/>Actuando en representation de mi familiar <b>{consentimiento.familiar_nombre}</b> (C.I. {consentimiento.familiar_cedula})."
        else:
            txt_testigo = ""
        texto = texto.replace("[BLOQUE_TESTIGO]", txt_testigo)

        parrafos = texto.split('\n')
        parrafos_filtrados = [p.strip() for p in parrafos if p.strip()]
        ultimo_parrafo_texto = parrafos_filtrados.pop() if parrafos_filtrados else ""

        for p in parrafos_filtrados:
            elementos.append(Paragraph(p, estilo_texto))

        style_firma_lbl = ParagraphStyle('FirmaLbl', fontName='Helvetica', fontSize=8, leading=10, alignment=1)
        
        if consentimiento.familiar_nombre:
            linea_paciente = Paragraph("_________________________________<br/><b>FIRMA DEL PACIENTE</b><br/>" + f"{paciente.nombre_completo}<br/>C.I: {paciente.cedula}", style_firma_lbl)
            linea_testigo = Paragraph("_________________________________<br/><b>FIRMA DEL FAMILIAR / TESTIGO</b><br/>" + f"{consentimiento.familiar_nombre}<br/>C.I: {consentimiento.familiar_cedula}", style_firma_lbl)
            celda_izquierda = [linea_paciente, Spacer(1, 15), linea_testigo]
        else:
            linea_paciente = Paragraph("_________________________________<br/><b>FIRMA DEL PACIENTE</b><br/>" + f"{paciente.nombre_completo}<br/>C.I: {paciente.cedula}", style_firma_lbl)
            celda_izquierda = [linea_paciente]

        linea_medico = Paragraph("_________________________________<br/><b>MÉDICO TRATANTE</b><br/>" + f"{prefijo} {doctor.nombre_doctor}<br/>MSDS: {doctor.matricula_msds}", style_firma_lbl)
        celda_derecha = [linea_medico]

        if doctor.firma_ruta and os.path.exists(doctor.firma_ruta):
            img_firma = Image(doctor.firma_ruta, width=1.6*inch, height=0.6*inch)
            img_firma.hAlign = 'CENTER'
            celda_derecha.insert(0, img_firma)
            if consentimiento.familiar_nombre:
                celda_izquierda.insert(0, Spacer(1, 0.4*inch))

        ancho_neto_util = (ancho_hoja / inch) - (80 / inch)
        datos_tabla = [[celda_izquierda, celda_derecha]]
        tabla_firmas = Table(datos_tabla, colWidths=[(ancho_neto_util/2)*inch, (ancho_neto_util/2)*inch])
        tabla_firmas.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('PADDING', (0,0), (-1,-1), 0),
        ]))

        bloque_inseparable = []
        if ultimo_parrafo_texto:
            bloque_inseparable.append(Paragraph(ultimo_parrafo_texto, estilo_texto))
        bloque_inseparable.append(Spacer(1, 12))
        bloque_inseparable.append(tabla_firmas)
        elementos.append(KeepTogether(bloque_inseparable))

        def footer_consentimiento(canvas, doc_obj):
            MotorImpresion._dibujar_cabecera(canvas, doctor, paciente, "CONSENTIMIENTO INFORMADO PARA CIRUGÍA")
            MotorImpresion._dibujar_pie_estatico_decorativo(canvas, doctor)
            canvas.setFont("Helvetica-Oblique", 8)
            canvas.setFillColor(colors.grey)

        try:
            doc.build(elementos, onFirstPage=footer_consentimiento, onLaterPages=footer_consentimiento)
            return os.path.abspath(ruta_salida_pdf)
        except Exception as e:
            raise Exception(f"Error al generar el PDF del consentimiento: {str(e)}")

    @staticmethod
    def generar_pdf_ordenes(doctor, paciente, lista_examenes, formato="pdf"):
        """Genera la orden de exámenes. Soporta formato="pdf" o formato="docx"."""
        if formato.lower() == "docx":
            ruta_base = os.path.abspath("recipes_generados")
            os.makedirs(ruta_base, exist_ok=True)
            nombre_base = f"OrdenExamen_{paciente.cedula}_{datetime.now().strftime('%Y-%m-%d')}"
            ruta_docx = MotorImpresion._obtener_ruta_unica(ruta_base, nombre_base, extension="docx")

            doc = MotorImpresion._crear_documento_docx()
            MotorImpresion._agregar_cabecera_docx(doc, doctor, paciente, "ORDEN DE EXÁMENES")

            p_head = doc.add_paragraph()
            p_head.paragraph_format.space_before = Pt(6)
            p_head.paragraph_format.space_after = Pt(8)
            r_head = p_head.add_run("Se solicita realizar:")
            r_head.bold = True
            r_head.font.size = Pt(12)

            for examen, obs in lista_examenes:
                p_ex = doc.add_paragraph()
                p_ex.paragraph_format.space_after = Pt(2)
                r_ex = p_ex.add_run(f"• {examen.strip().title()}")
                r_ex.bold = True
                r_ex.font.size = Pt(11)

                if obs.strip():
                    p_obs = doc.add_paragraph()
                    p_obs.paragraph_format.left_indent = Inches(0.3)
                    p_obs.paragraph_format.space_after = Pt(6)
                    r_obs = p_obs.add_run(f"Obs: {obs.strip()}")
                    r_obs.font.italic = True
                    r_obs.font.size = Pt(9.5)
                    r_obs.font.color.rgb = RGBColor(70, 80, 95)

            MotorImpresion._agregar_firma_docx(doc, doctor)
            MotorImpresion._agregar_pie_pagina_docx(doc, doctor)
            doc.save(ruta_docx)
            return os.path.abspath(ruta_docx)

        # --- GENERACIÓN PDF ORIGINAL ---
        ruta_base = os.path.abspath("recipes_generados")
        os.makedirs(ruta_base, exist_ok=True)
        
        nombre_base = f"OrdenExamen_{paciente.cedula}_{datetime.now().strftime('%Y-%m-%d')}"
        ruta_completa = MotorImpresion._obtener_ruta_unica(ruta_base, nombre_base, extension="pdf")
        
        doc = SimpleDocTemplate(
            ruta_completa, 
            pagesize=A4, 
            leftMargin=0.5*inch, 
            rightMargin=0.5*inch,
            topMargin=2.5*inch,    
            bottomMargin=0.9*inch  
        )
        
        styles = getSampleStyleSheet()
        estilo_examen = ParagraphStyle('Examen', fontName='Helvetica-Bold', fontSize=11, leading=14, spaceBefore=6)
        estilo_obs = ParagraphStyle('Obs', fontName='Helvetica-Oblique', fontSize=10, leading=14, leftIndent=15, spaceAfter=6, textColor=colors.HexColor("#475569"))
        
        elementos = []
        elementos.append(Paragraph("<font size=14><b>Se solicita realizar:</b></font>", estilo_examen))
        elementos.append(Spacer(1, 10))
        
        for examen, obs in lista_examenes:
            bloque_examen = []
            ex_formateado = examen.strip().title()
            
            p_ex = Paragraph(f"• {ex_formateado}", estilo_examen)
            bloque_examen.append(p_ex)
            
            if obs.strip():
                p_obs = Paragraph(f"Obs: {obs.strip()}", estilo_obs)
                bloque_examen.append(p_obs)
                
            elementos.append(KeepTogether(bloque_examen))

        elementos.append(Spacer(1, 25))
        elementos.append(KeepTogether([MotorImpresion._obtener_bloque_firma_fluyente(doctor)]))

        def construir_paginas_ordenes(canvas_obj, doc_obj):
            MotorImpresion._dibujar_cabecera(canvas_obj, doctor, paciente, "ORDEN DE EXÁMENES")
            MotorImpresion._dibujar_pie_estatico_decorativo(canvas_obj, doctor)

        try:
            doc.build(elementos, onFirstPage=construir_paginas_ordenes, onLaterPages=construir_paginas_ordenes)
            return os.path.abspath(ruta_completa)
        except Exception as e:
            raise Exception(f"Error al generar el PDF de órdenes: {str(e)}")

    @staticmethod
    def generar_pdf_egreso(doctor, paciente, egreso, ruta_salida, formato="pdf"):
        """Genera Informe de Egreso. Soporta formato="pdf" o formato="docx"."""
        if formato.lower() == "docx":
            ruta_docx = MotorImpresion._ajustar_extension_salida(ruta_salida, "docx")
            doc = MotorImpresion._crear_documento_docx()
            MotorImpresion._agregar_cabecera_docx(doc, doctor, paciente, "INFORME DE EGRESO")

            def add_sec(tit, cnt):
                if cnt and str(cnt).strip().lower() not in ["no aplica", "none", "", "n/a"]:
                    p_t = doc.add_paragraph()
                    p_t.paragraph_format.space_before = Pt(8)
                    p_t.paragraph_format.space_after = Pt(2)
                    r = p_t.add_run(tit)
                    r.bold = True
                    r.font.size = Pt(11)
                    r.font.color.rgb = RGBColor(142, 68, 173)

                    p_c = doc.add_paragraph()
                    p_c.paragraph_format.space_after = Pt(6)
                    p_c.add_run(str(cnt))

            add_sec("Resumen del Egreso", egreso.resumen_egreso)
            add_sec("Evaluación", egreso.evaluacion)
            add_sec("Examen Físico", egreso.examen_fisico)
            add_sec("Paraclínicos", egreso.paraclinicos)
            add_sec("Diagnóstico de Egreso", egreso.diagnostico_egreso)
            add_sec("Procedimiento Realizado", egreso.procedimiento_realizado)
            add_sec("Plan de Manejo", egreso.plan_manejo)

            MotorImpresion._agregar_firma_docx(doc, doctor)
            MotorImpresion._agregar_pie_pagina_docx(doc, doctor)
            doc.save(ruta_docx)
            return os.path.abspath(ruta_docx)

        # --- GENERACIÓN PDF ORIGINAL ---
        ruta_salida_pdf = MotorImpresion._ajustar_extension_salida(ruta_salida, "pdf")
        doc = SimpleDocTemplate(ruta_salida_pdf, pagesize=A4, topMargin=2.5*inch, bottomMargin=2.0*inch, leftMargin=1*inch, rightMargin=1*inch)
        styles = getSampleStyleSheet()
        estilo_seccion = ParagraphStyle('Seccion', parent=styles['Heading3'], color=colors.HexColor("#8e44ad"), spaceBefore=12, spaceAfter=6)
        estilo_texto = ParagraphStyle('Texto', parent=styles['Normal'], fontSize=11, leading=15, alignment=4) 
        
        elementos = []

        def agregar_seccion(titulo, contenido):
            if contenido and str(contenido).strip().lower() not in ["no aplica", "none", "", "n/a"]:
                elementos.append(Paragraph(f"<b>{titulo}</b>", estilo_seccion))
                texto_formateado = str(contenido).replace('\n', '<br/>')
                elementos.append(Paragraph(texto_formateado, estilo_texto))
                elementos.append(Spacer(1, 10))

        agregar_seccion("Resumen del Egreso", egreso.resumen_egreso)
        agregar_seccion("Evaluación", egreso.evaluacion)
        agregar_seccion("Examen Físico", egreso.examen_fisico)
        agregar_seccion("Paraclínicos", egreso.paraclinicos)
        agregar_seccion("Diagnóstico de Egreso", egreso.diagnostico_egreso)
        agregar_seccion("Procedimiento Realizado", egreso.procedimiento_realizado)
        agregar_seccion("Plan de Manejo", egreso.plan_manejo)

        if not elementos:
            elementos.append(Paragraph("Informe sin detalles clínicos registrados.", estilo_texto))

        elementos.append(Spacer(1, 20))
        elementos.append(KeepTogether([MotorImpresion._obtener_bloque_firma_fluyente(doctor)]))

        def footer_egreso(canvas, doc_obj):
            MotorImpresion._dibujar_cabecera(canvas, doctor, paciente, "INFORME DE EGRESO")
            canvas.setFont("Helvetica-Oblique", 8)
            canvas.setFillColor(colors.grey)
            MotorImpresion._dibujar_pie_estatico_decorativo(canvas, doctor)

        doc.build(elementos, onFirstPage=footer_egreso, onLaterPages=footer_egreso)
        return ruta_salida_pdf

    @staticmethod
    def generar_pdf_postop(doctor, paciente, nota, ruta_salida, formato="pdf"):
        """Genera Informe Post-operatorio. Soporta formato="pdf" o formato="docx"."""
        if formato.lower() == "docx":
            ruta_docx = MotorImpresion._ajustar_extension_salida(ruta_salida, "docx")
            doc = MotorImpresion._crear_documento_docx()
            MotorImpresion._agregar_cabecera_docx(doc, doctor, paciente, "INFORME POST-OPERATORIO")

            def add_sec(tit, cnt):
                if cnt and str(cnt).strip().lower() not in ["no aplica", "none", "", "n/a"]:
                    p_t = doc.add_paragraph()
                    p_t.paragraph_format.space_before = Pt(8)
                    p_t.paragraph_format.space_after = Pt(2)
                    r = p_t.add_run(tit)
                    r.bold = True
                    r.font.size = Pt(11)
                    r.font.color.rgb = RGBColor(192, 57, 43)

                    p_c = doc.add_paragraph()
                    p_c.paragraph_format.space_after = Pt(6)
                    p_c.add_run(str(cnt))

            add_sec("Diagnóstico Pre-operatorio", nota.diagnostico_preoperatorio)
            add_sec("Cirugía Realizada", nota.cirugia_realizada)
            add_sec("Diagnóstico Post-operatorio", nota.diagnostico_postoperatorio)
            add_sec("Descripción de la intervención", nota.descripcion_intervencion)
            add_sec("Hallazgos", nota.hallazgos)
            add_sec("Biopsia", nota.biopsia)
            add_sec("Plan de Manejo", nota.plan_manejo)

            MotorImpresion._agregar_firma_docx(doc, doctor)
            MotorImpresion._agregar_pie_pagina_docx(doc, doctor)
            doc.save(ruta_docx)
            return os.path.abspath(ruta_docx)

        # --- GENERACIÓN PDF ORIGINAL ---
        ruta_salida_pdf = MotorImpresion._ajustar_extension_salida(ruta_salida, "pdf")
        doc = SimpleDocTemplate(ruta_salida_pdf, pagesize=A4, topMargin=2.5*inch, bottomMargin=2.0*inch, leftMargin=1*inch, rightMargin=1*inch)
        styles = getSampleStyleSheet()
        estilo_seccion = ParagraphStyle('Seccion', parent=styles['Heading3'], color=colors.HexColor("#c0392b"), spaceBefore=12, spaceAfter=6)
        estilo_texto = ParagraphStyle('Texto', parent=styles['Normal'], fontSize=11, leading=15, alignment=4) 
        
        elementos = []

        def agregar_seccion(titulo, contenido):
            if contenido and str(contenido).strip().lower() not in ["no aplica", "none", "", "n/a"]:
                elementos.append(Paragraph(f"<b>{titulo}</b>", estilo_seccion))
                texto_formateado = str(contenido).replace('\n', '<br/>')
                elementos.append(Paragraph(texto_formateado, estilo_texto))
                elementos.append(Spacer(1, 10))

        agregar_seccion("Diagnóstico Pre-operatorio", nota.diagnostico_preoperatorio)
        agregar_seccion("Cirugía Realizada", nota.cirugia_realizada)
        agregar_seccion("Diagnóstico Post-operatorio", nota.diagnostico_postoperatorio)
        agregar_seccion("Descripción detallada de la intervención", nota.descripcion_intervencion)
        agregar_seccion("Hallazgos", nota.hallazgos)
        agregar_seccion("Biopsia", nota.biopsia)
        agregar_seccion("Plan de Manejo", nota.plan_manejo)

        if not elementos:
            elementos.append(Paragraph("Informe sin detalles quirúrgicos registrados.", estilo_texto))

        elementos.append(Spacer(1, 20))
        elementos.append(KeepTogether([MotorImpresion._obtener_bloque_firma_fluyente(doctor)]))

        def footer_postop(canvas, doc_obj):
            MotorImpresion._dibujar_cabecera(canvas, doctor, paciente, "INFORME POST-OPERATORIO")
            canvas.setFont("Helvetica-Oblique", 8)
            canvas.setFillColor(colors.grey)
            MotorImpresion._dibujar_pie_estatico_decorativo(canvas, doctor)

        doc.build(elementos, onFirstPage=footer_postop, onLaterPages=footer_postop)
        return ruta_salida_pdf

    @staticmethod
    def generar_pdf_factura(doctor, datos, conceptos, ruta_salida, formato="pdf"):
        """Genera Factura. Soporta formato="pdf" o formato="docx"."""
        if formato.lower() == "docx":
            ruta_docx = MotorImpresion._ajustar_extension_salida(ruta_salida, "docx")
            doc = MotorImpresion._crear_documento_docx()
            MotorImpresion._agregar_cabecera_docx(doc, doctor, {"nombre_completo": datos["cliente_nombre"], "cedula": datos["cliente_cedula"], "edad_actual": "--"}, f"FACTURA N° {datos['num_factura']}")

            p_meta = doc.add_paragraph()
            p_meta.add_run(f"Nº Control: {datos['num_control']}   |   Emisión: {datos['fecha_emision']}   |   Condición: {datos['condiciones_pago']}\n")
            p_meta.add_run(f"Cliente: {datos['cliente_nombre']} (RIF/CI: {datos['cliente_cedula']})\n")
            p_meta.add_run(f"Domicilio: {datos['cliente_domicilio']}")

            table = doc.add_table(rows=1, cols=4)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Cant.'
            hdr_cells[1].text = 'Descripción'
            hdr_cells[2].text = f'P.Unit ({datos["moneda"]})'
            hdr_cells[3].text = f'Total ({datos["moneda"]})'

            for item in conceptos:
                if item["descripcion"].strip():
                    row_cells = table.add_row().cells
                    row_cells[0].text = str(item["cantidad"])
                    row_cells[1].text = str(item["descripcion"])
                    row_cells[2].text = str(item["precio_unitario"])
                    row_cells[3].text = str(item["total_renglon"])

            p_tot = doc.add_paragraph()
            p_tot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_tot.paragraph_format.space_before = Pt(12)
            r_t = p_tot.add_run(f"TOTAL EXENTO: {datos['total_pagar']} {datos['moneda']}")
            r_t.bold = True
            r_t.font.size = Pt(14)

            MotorImpresion._agregar_pie_pagina_docx(doc, doctor)
            doc.save(ruta_docx)
            return os.path.abspath(ruta_docx)

        # --- GENERACIÓN PDF ORIGINAL ---
        ruta_salida_pdf = MotorImpresion._ajustar_extension_salida(ruta_salida, "pdf")
        ancho_hoja, alto_hoja = A4
        doc = SimpleDocTemplate(ruta_salida_pdf, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()
        
        style_clinica = ParagraphStyle('TxtClinica', fontName='Helvetica-Bold', fontSize=12, leading=14, alignment=1)
        style_doc_title = ParagraphStyle('TxtDocTitle', fontName='Times-BoldItalic', fontSize=22, leading=24, alignment=1)
        style_meta_doc = ParagraphStyle('TxtMetaDoc', fontName='Helvetica', fontSize=10, leading=13, alignment=1, textColor=colors.HexColor("#334155"))
        style_rif = ParagraphStyle('TxtRif', fontName='Helvetica-Bold', fontSize=12, leading=14, alignment=1)
        style_legal_num = ParagraphStyle('TxtLegalNum', fontName='Helvetica-Bold', fontSize=11, leading=14, alignment=1)
        style_client_lbl = ParagraphStyle('TxtClientLbl', fontName='Helvetica-Bold', fontSize=10, leading=12)
        style_client_val = ParagraphStyle('TxtClientVal', fontName='Helvetica', fontSize=10, leading=12)
        
        style_th = ParagraphStyle('TxtTH', fontName='Helvetica-Bold', fontSize=10, leading=14, alignment=1, textColor=colors.white)
        style_td = ParagraphStyle('TxtTD', fontName='Helvetica', fontSize=10, leading=13)
        style_td_center = ParagraphStyle('TxtTDCenter', fontName='Helvetica', fontSize=10, leading=13, alignment=1)
        style_td_right = ParagraphStyle('TxtTDRight', fontName='Helvetica-Bold', fontSize=10, leading=13, alignment=2)
        
        elementos = []
        moneda = datos["moneda"]

        es_fem = hasattr(doctor, 'genero_doctor') and doctor.genero_doctor and doctor.genero_doctor.strip().lower() in ["femenino", "mujer"]
        prefijo = "Dra." if es_fem else "Dr."
        
        clinica_txt = f"<b>{doctor.nombre_clinica.upper()}</b>" if doctor.nombre_clinica else "<b>CENTRO MÉDICO</b>"
        col_logo = [Paragraph(clinica_txt, style_clinica)]
        
        if doctor.logo_ruta and os.path.exists(doctor.logo_ruta):
            img_logo = Image(doctor.logo_ruta, width=1.1*inch, height=0.8*inch)
            img_logo.hAlign = 'CENTER'
            col_logo.insert(0, img_logo)
            
        col_centro = [
            Paragraph(f"{prefijo} {doctor.nombre_doctor.upper()}", style_doc_title),
            Paragraph(doctor.especialidades.upper() if doctor.especialidades else "ESPECIALISTA", style_meta_doc),
            Paragraph(doctor.direccion_clinica if doctor.direccion_clinica else "Dirección no registrada", style_meta_doc),
            Paragraph(f"C.I: {doctor.cedula_doctor or '--'} | MSDS: {doctor.matricula_msds or '--'}", style_meta_doc),
            Paragraph(f"Contacto: {doctor.correo_doctor or '--'}", style_meta_doc)
        ]
        
        col_derecha = [
            Paragraph(f"<b>RIF: {doctor.rif_clinica.upper() if doctor.rif_clinica else '--'}</b>", style_rif),
            Spacer(1, 10),
            Paragraph(f"<b>FECHA DE EMISIÓN</b><br/>{datos['fecha_emision']}", style_meta_doc)
        ]
        
        ancho_util_factura = ancho_hoja - 72
        tabla_membrete = Table([[col_logo, col_centro, col_derecha]], colWidths=[ancho_util_factura*0.18, ancho_util_factura*0.64, ancho_util_factura*0.18])
        tabla_membrete.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 10),
        ]))
        elementos.append(tabla_membrete)
        elementos.append(Spacer(1, 10))

        txt_control = Paragraph(f"<b>Nº de Control:</b> {datos['num_control']}", style_legal_num)
        txt_factura = Paragraph(f"<font color='#e11d48'><b>Factura Nº:</b> {datos['num_factura']}</font>", style_legal_num)
        
        tabla_num_legal = Table([[txt_control, txt_factura]], colWidths=[ancho_util_factura*0.5, ancho_util_factura*0.5])
        tabla_num_legal.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f8fafc")),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#cbd5e1")),
            ('PADDING', (0,0), (-1,-1), 8),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE')
        ]))
        elementos.append(tabla_num_legal)
        elementos.append(Spacer(1, 12))

        grid_datos_cliente = [
            [Paragraph("Nombre / Razón Social:", style_client_lbl), Paragraph(datos["cliente_nombre"], style_client_val), "", "", "", ""],
            [Paragraph("Domicilio Fiscal:", style_client_lbl), Paragraph(datos["cliente_domicilio"], style_client_val), "", "", "", ""],
            [
                Paragraph("RIF / C.I.:", style_client_lbl), Paragraph(datos["cliente_cedula"], style_client_val),
                Paragraph("Teléfono:", style_client_lbl), Paragraph(datos["cliente_telefono"], style_client_val),
                Paragraph("Condición:", style_client_lbl), Paragraph(datos["condiciones_pago"], style_client_val)
            ]
        ]
        
        tabla_cliente = Table(grid_datos_cliente, colWidths=[ancho_util_factura*0.20, ancho_util_factura*0.24, ancho_util_factura*0.11, ancho_util_factura*0.16, ancho_util_factura*0.12, ancho_util_factura*0.17])
        tabla_cliente.setStyle(TableStyle([
            ('SPAN', (1,0), (5,0)), 
            ('SPAN', (1,1), (5,1)), 
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#cbd5e1")),
            ('PADDING', (0,0), (-1,-1), 6),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('BACKGROUND', (0,0), (-1,-1), colors.white)
        ]))
        elementos.append(tabla_cliente)
        elementos.append(Spacer(1, 15))

        headers_tabla = [
            Paragraph("Cant.", style_th), 
            Paragraph("Descripción del Servicio / Concepto", style_th), 
            Paragraph(f"Precio Unitario ({moneda})", style_th),
            Paragraph(f"Total ({moneda})", style_th)
        ]
        
        filas_conceptos = [headers_tabla]
        for item in conceptos:
            if item["descripcion"].strip():
                filas_conceptos.append([
                    Paragraph(item["cantidad"], style_td_center),
                    Paragraph(item["descripcion"], style_td),
                    Paragraph(item["precio_unitario"], style_td_right),
                    Paragraph(item["total_renglon"], style_td_right)
                ])

        tabla_factura_conceptos = Table(filas_conceptos, colWidths=[ancho_util_factura*0.08, ancho_util_factura*0.60, ancho_util_factura*0.16, ancho_util_factura*0.16])
        tabla_factura_conceptos.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1e293b")), 
            ('GRID', (0,0), (-1,-1), 1, colors.HexColor("#cbd5e1")),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('PADDING', (0,0), (-1,0), 8),
            ('PADDING', (0,1), (-1,-1), 6),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#f8fafc")])
        ]))
        elementos.append(tabla_factura_conceptos)
        elementos.append(Spacer(1, 25))

        lbl_pago = [
            Paragraph(f"<b>Forma de Pago:</b> {datos['forma_pago']}", style_client_lbl),
            Spacer(1, 5),
            Paragraph("<font color='#64748b'>Documento Emitido conforme a las providencias del SENIAT para Servicios Médicos Exentos.</font>", style_meta_doc)
        ]
        
        cuadro_sello = Table([
            [Paragraph("<b>SELLO</b>", style_td_center)],
            [Spacer(1, 30)],
            [Paragraph("<small>RECIBÍ CONFORME</small>", style_td_center)]
        ], colWidths=[ancho_util_factura*0.28])
        cuadro_sello.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 1.5, colors.HexColor("#94a3b8")),
            ('LINESTYLE', (0,0), (-1,-1), 'DASHED'),
            ('PADDING', (0,0), (-1,-1), 4),
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f8fafc"))
        ]))
        
        style_total_exento = ParagraphStyle('TxtTotalExento', fontName='Helvetica-Bold', fontSize=16, leading=25, alignment=2)
        bloque_totales = [
            Paragraph(f"<font size=11><b>MONEDA:</b> {combo_moneda_txt(moneda)}</font>", style_client_lbl),
            Spacer(1, 15),
            Paragraph(f"<font size=14 color='#0f172a'><b>TOTAL EXENTO:</b></font><br/>"
                      f"<font size=18 color='#16a34a'><b>{datos['total_pagar']}</b></font>", style_total_exento)
        ]
        
        tabla_pie = Table([[lbl_pago, cuadro_sello, bloque_totales]], colWidths=[ancho_util_factura*0.38, ancho_util_factura*0.30, ancho_util_factura*0.32])
        tabla_pie.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('PADDING', (0,0), (-1,-1), 5)
        ]))
        
        elementos.append(KeepTogether([tabla_pie]))

        try:
            doc.build(elementos)
            return os.path.abspath(ruta_salida_pdf)
        except Exception as e:
            raise Exception(f"Error crítico estructurando la Factura PDF: {str(e)}")

    @staticmethod
    def limpiar_carpeta_pdf():
        if os.path.exists(MotorImpresion.CARPETA_RECIPES):
            try:
                for archivo in os.listdir(MotorImpresion.CARPETA_RECIPES):
                    ruta_archivo = os.path.join(MotorImpresion.CARPETA_RECIPES, archivo)
                    if os.path.isfile(ruta_archivo):
                        os.remove(ruta_archivo)
                return True
            except:
                return False
        return True

    @staticmethod
    def abrir_documento_generado(ruta_archivo):
        """Abre la ruta pasada (soporta .pdf y .docx automáticamente)."""
        if not os.path.exists(ruta_archivo):
            base, _ = os.path.splitext(ruta_archivo)
            if os.path.exists(f"{base}.docx"):
                ruta_archivo = f"{base}.docx"
            elif os.path.exists(f"{base}.pdf"):
                ruta_archivo = f"{base}.pdf"

        if os.path.exists(ruta_archivo):
            from PySide6.QtGui import QDesktopServices
            from PySide6.QtCore import QUrl
            QDesktopServices.openUrl(QUrl.fromLocalFile(os.path.abspath(ruta_archivo)))
            return True
        return False


def combo_moneda_txt(simbolo):
    return "DIVISA (USD $)" if "$" in simbolo else "BOLÍVARES (Bs.)"